#!/usr/bin/env python3
"""Genere un fichier Excel de leads a partir de Google Maps (Places API New).

Criteres par defaut (config.json) :
  - note Google >= 4.7
  - au moins N avis (evite les 5.0 avec 1 avis)
  - AUCUN site web renseigne sur la fiche  <- la cible commerciale
  - un numero de telephone disponible
  - etablissement en activite

Usage :
    export GOOGLE_MAPS_API_KEY="..."
    python3 generate_leads.py --config config.json --out leads_maps.xlsx
"""

import argparse
import json
import os
import sys
import time
from typing import Dict, Iterator, List

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SEARCH_URL = "https://places.googleapis.com/v1/places:searchText"

FIELD_MASK = ",".join(
    [
        "nextPageToken",
        "places.id",
        "places.displayName",
        "places.formattedAddress",
        "places.nationalPhoneNumber",
        "places.internationalPhoneNumber",
        "places.websiteUri",
        "places.rating",
        "places.userRatingCount",
        "places.googleMapsUri",
        "places.primaryTypeDisplayName",
        "places.businessStatus",
    ]
)

COLUMNS = [
    ("Categorie", 22),
    ("Nom", 38),
    ("Note", 8),
    ("Avis", 8),
    ("Telephone", 18),
    ("Telephone int.", 20),
    ("Adresse", 52),
    ("Type", 22),
    ("Fiche Google Maps", 46),
    ("Statut contact", 16),
    ("Notes", 30),
]


def search_category(api_key: str, query: str, language: str, region: str,
                    max_pages: int = 3) -> Iterator[Dict]:
    """Itere sur les resultats Text Search (20 par page, 3 pages max cote Google)."""
    headers = {
        "Content-Type": "application/json",
        "X-Goog-Api-Key": api_key,
        "X-Goog-FieldMask": FIELD_MASK,
    }
    body = {"textQuery": query, "languageCode": language, "regionCode": region}
    for page in range(max_pages):
        resp = requests.post(SEARCH_URL, headers=headers, json=body, timeout=30)
        if resp.status_code != 200:
            print(f"  ! API {resp.status_code}: {resp.text[:300]}", file=sys.stderr)
            return
        data = resp.json()
        for place in data.get("places", []):
            yield place
        token = data.get("nextPageToken")
        if not token:
            return
        body = {"pageToken": token}
        time.sleep(2)  # le token n'est pas actif immediatement


def keep(place: Dict, f: Dict) -> bool:
    rating = place.get("rating")
    if rating is None:
        return False
    if f.get("strict_greater_than"):
        if rating <= f["min_rating"]:
            return False
    elif rating < f["min_rating"]:
        return False
    if place.get("userRatingCount", 0) < f.get("min_reviews", 0):
        return False
    if f.get("require_no_website") and place.get("websiteUri"):
        return False
    if f.get("require_phone") and not place.get("nationalPhoneNumber"):
        return False
    if f.get("operational_only") and place.get("businessStatus", "OPERATIONAL") != "OPERATIONAL":
        return False
    return True


def collect(cfg: Dict, api_key: str) -> List[Dict]:
    filters = cfg["filters"]
    seen, rows = set(), []
    for category in cfg["categories"]:
        query = f"{category} {cfg['city']}"
        kept = 0
        for place in search_category(api_key, query, cfg["language"], cfg["region"]):
            pid = place.get("id")
            if not pid or pid in seen:
                continue
            if not keep(place, filters):
                continue
            seen.add(pid)
            kept += 1
            rows.append(
                {
                    "Categorie": category,
                    "Nom": place.get("displayName", {}).get("text", ""),
                    "Note": place.get("rating"),
                    "Avis": place.get("userRatingCount"),
                    "Telephone": place.get("nationalPhoneNumber", ""),
                    "Telephone int.": place.get("internationalPhoneNumber", ""),
                    "Adresse": place.get("formattedAddress", ""),
                    "Type": place.get("primaryTypeDisplayName", {}).get("text", ""),
                    "Fiche Google Maps": place.get("googleMapsUri", ""),
                    "Statut contact": "A contacter",
                    "Notes": "",
                }
            )
        print(f"  {category:<28} {kept} lead(s) retenu(s)")
    rows.sort(key=lambda r: (r["Categorie"], -(r["Note"] or 0), -(r["Avis"] or 0)))
    return rows


def write_xlsx(rows: List[Dict], path: str, cfg: Dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    ws.append([c[0] for c in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width

    for r in rows:
        ws.append([r[c[0]] for c in COLUMNS])

    link_col = get_column_letter([c[0] for c in COLUMNS].index("Fiche Google Maps") + 1)
    for row_idx in range(2, ws.max_row + 1):
        cell = ws[f"{link_col}{row_idx}"]
        if cell.value:
            cell.hyperlink = cell.value
            cell.value = "Voir la fiche"
            cell.font = Font(color="0563C1", underline="single")

    if ws.max_row > 1:
        ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"
        table = Table(displayName="Leads", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(table)
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Criteres")
    f = cfg["filters"]
    op = ">" if f.get("strict_greater_than") else ">="
    for line in [
        ("Ville", cfg["city"]),
        ("Note Google", f"{op} {f['min_rating']}"),
        ("Nombre d'avis minimum", f["min_reviews"]),
        ("Site web", "aucun site web sur la fiche" if f.get("require_no_website") else "indifferent"),
        ("Telephone", "obligatoire" if f.get("require_phone") else "indifferent"),
        ("Etablissements", "en activite uniquement" if f.get("operational_only") else "tous"),
        ("Categories", ", ".join(cfg["categories"])),
        ("Total leads", len(rows)),
        ("Source", "Google Maps - Places API (New), Text Search"),
    ]:
        info.append(list(line))
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 90
    for row in info.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    info["B7"].alignment = Alignment(wrap_text=True)

    wb.save(path)


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--config", default=os.path.join(os.path.dirname(__file__), "config.json"))
    p.add_argument("--out")
    p.add_argument("--city", help="surcharge la ville du config")
    args = p.parse_args()

    api_key = os.environ.get("GOOGLE_MAPS_API_KEY")
    if not api_key:
        print("GOOGLE_MAPS_API_KEY manquante (Places API New a activer sur le projet).",
              file=sys.stderr)
        return 2

    with open(args.config, encoding="utf-8") as fh:
        cfg = json.load(fh)
    if args.city:
        cfg["city"] = args.city
    out = args.out or cfg.get("output", "leads_maps.xlsx")

    print(f"Recherche Google Maps - {cfg['city']}")
    rows = collect(cfg, api_key)
    if not rows:
        print("Aucun lead ne correspond aux criteres.", file=sys.stderr)
    write_xlsx(rows, out, cfg)
    print(f"\n{len(rows)} lead(s) -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
